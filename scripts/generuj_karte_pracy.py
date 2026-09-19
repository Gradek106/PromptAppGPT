#!/usr/bin/env python3
"""
Generator Kart Pracy — samodzielny skrypt produkcyjny.

Generuje gotowe do druku karty pracy dla szkoły podstawowej (klasy 1-8)
wraz z opcjonalnym kluczem odpowiedzi dla nauczyciela, wywołując
OpenAI API i zapisując wynik jako plik .docx (Word) oraz .txt.

Wykorzystuje tę samą logikę promptów co aplikacja PromptAppGPT
`app/karty_pracy.yml`, ale działa jako niezależny skrypt CLI —
bez potrzeby uruchamiania interfejsu webowego.

Wymagania:
    pip install -r requirements.txt

Klucz API:
    Ustaw zmienną środowiskową OPENAI_API_KEY, albo przekaż --api-key.

Przykłady użycia
----------------

Pojedyncza karta pracy:

    python scripts/generuj_karte_pracy.py \\
        --przedmiot "Matematyka" --klasa 4 --temat "Mnożenie pisemne" \\
        --liczba-zadan 5 --trudnosc "zróżnicowany" --klucz tak \\
        --output karty/mnozenie_pisemne

Produkcja wsadowa (wiele kart naraz) z pliku CSV:

    python scripts/generuj_karte_pracy.py --batch tematy.csv --output-dir karty/

    Plik CSV powinien mieć kolumny:
    przedmiot,klasa,temat,liczba_zadan,trudnosc,typ_zadan,klucz,wskazowki
"""

from __future__ import annotations

import argparse
import csv
import os
import re
import sys
import time
import unicodedata
from dataclasses import dataclass, fields
from pathlib import Path
from typing import Optional

try:
    from openai import OpenAI
except ImportError:  # pragma: no cover
    print(
        "Brak pakietu 'openai'. Zainstaluj zależności poleceniem:\n"
        "    pip install -r requirements.txt",
        file=sys.stderr,
    )
    sys.exit(1)

try:
    from docx import Document
    from docx.shared import Pt
except ImportError:  # pragma: no cover
    print(
        "Brak pakietu 'python-docx'. Zainstaluj zależności poleceniem:\n"
        "    pip install -r requirements.txt",
        file=sys.stderr,
    )
    sys.exit(1)


SYSTEM_PROMPT = """\
Jesteś doświadczonym nauczycielem szkoły podstawowej i metodykiem nauczania.
Tworzysz karty pracy zgodne z polską podstawą programową.
Zawsze piszesz po polsku, poprawnie, prostym i przyjaznym dla ucznia językiem.
Dostosowujesz słownictwo, długość poleceń i trudność zadań do wieku uczniów danej klasy.
Odpowiadasz wyłącznie treścią karty pracy, bez komentarzy i wstępów."""

USER_PROMPT_TEMPLATE = """\
Przygotuj kartę pracy dla szkoły podstawowej według poniższych parametrów.

Przedmiot: {przedmiot}
Klasa: {klasa}
Temat: {temat}
Cele lekcji lub dodatkowe wskazówki (opcjonalnie): {wskazowki}
Liczba zadań: {liczba_zadan}
Poziom trudności: {trudnosc}
Typy zadań: {typ_zadan}
Klucz odpowiedzi dla nauczyciela: {klucz}

Wymagania dotyczące karty pracy:
1. Nagłówek z tytułem tematu, przedmiotem, klasą oraz miejscem na imię, nazwisko i datę (Imię i nazwisko: ............ Data: ............).
2. Krótkie, jednozdaniowe wprowadzenie do tematu napisane językiem zrozumiałym dla ucznia tej klasy.
3. Dokładnie tyle zadań, ile podano w parametrach, ponumerowanych "Zadanie 1", "Zadanie 2" itd.
4. Każde zadanie ma jasne polecenie, a pod nim miejsce na odpowiedź (linie z kropek lub pola do uzupełnienia) oraz liczbę punktów w nawiasie.
5. Zadania mają być zgodne z podstawą programową dla podanej klasy i realizować podany temat.
6. Jeśli wybrano poziom "zróżnicowany", uporządkuj zadania od najłatwiejszego do najtrudniejszego, a ostatnie oznacz jako "Zadanie dla chętnych".
7. Na końcu dodaj krótką sekcję "Samoocena" z 3 zdaniami do zaznaczenia (np. "Rozumiem temat", "Potrzebuję jeszcze poćwiczyć", "Chcę zapytać nauczyciela").
8. Jeśli klucz odpowiedzi = tak, po karcie pracy dodaj oddzieloną linią "----------" sekcję "KLUCZ ODPOWIEDZI (dla nauczyciela)" zawierającą odpowiedzi do wszystkich zadań, punktację oraz łączną liczbę punktów. Jeśli klucz odpowiedzi = nie, pomiń tę sekcję.
9. Nie używaj formatowania Markdown (bez gwiazdek i krzyżyków), tylko zwykły tekst gotowy do skopiowania i wydrukowania.
"""

ANSWER_KEY_MARKER = "KLUCZ ODPOWIEDZI"


@dataclass
class KartaParams:
    przedmiot: str = "Matematyka"
    klasa: str = "4"
    temat: str = ""
    wskazowki: str = "brak"
    liczba_zadan: str = "5"
    trudnosc: str = "zróżnicowany (od łatwych do trudnych)"
    typ_zadan: str = "mieszane"
    klucz: str = "tak"

    def prompt(self) -> str:
        wskazowki = self.wskazowki.strip() or "brak"
        return USER_PROMPT_TEMPLATE.format(
            przedmiot=self.przedmiot,
            klasa=self.klasa,
            temat=self.temat,
            wskazowki=wskazowki,
            liczba_zadan=self.liczba_zadan,
            trudnosc=self.trudnosc,
            typ_zadan=self.typ_zadan,
            klucz=self.klucz,
        )


def slugify(text: str) -> str:
    text = unicodedata.normalize("NFKD", text).encode("ascii", "ignore").decode("ascii")
    text = re.sub(r"[^a-zA-Z0-9]+", "_", text).strip("_").lower()
    return text or "karta_pracy"


def generate_worksheet_text(
    client: OpenAI,
    params: KartaParams,
    model: str,
    max_retries: int = 2,
) -> str:
    """Wywołuje OpenAI API i zwraca tekst karty pracy, z prostą walidacją i ponowieniami."""
    last_error: Optional[Exception] = None
    for attempt in range(max_retries + 1):
        try:
            response = client.chat.completions.create(
                model=model,
                messages=[
                    {"role": "system", "content": SYSTEM_PROMPT},
                    {"role": "user", "content": params.prompt()},
                ],
                temperature=0.7,
            )
            text = (response.choices[0].message.content or "").strip()
            if "Zadanie 1" in text and len(text) >= 100:
                return text
            last_error = ValueError(
                "Odpowiedź nie przeszła walidacji (brak 'Zadanie 1' lub zbyt krótka)."
            )
        except Exception as exc:  # noqa: BLE001 - chcemy złapać i zalogować, potem ponowić
            last_error = exc
        if attempt < max_retries:
            time.sleep(2 * (attempt + 1))
    raise RuntimeError(
        f"Nie udało się wygenerować karty pracy dla tematu '{params.temat}': {last_error}"
    )


def split_worksheet_and_key(text: str) -> tuple[str, Optional[str]]:
    if ANSWER_KEY_MARKER in text:
        idx = text.index(ANSWER_KEY_MARKER)
        before = text[:idx]
        divider_idx = before.rfind("----------")
        worksheet = (before[:divider_idx] if divider_idx != -1 else before).rstrip()
        key = text[idx:].strip()
        return worksheet, key
    return text.strip(), None


def save_as_docx(text: str, params: KartaParams, path: Path) -> None:
    """Zapisuje kartę pracy jako .docx, formatując dokładnie to, co zwrócił model
    (bez dodawania własnego nagłówka, żeby uniknąć zdublowania tytułu/przedmiotu,
    które GPT już umieszcza na początku tekstu, zgodnie z promptem)."""
    worksheet_text, key_text = split_worksheet_and_key(text)

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    lines = worksheet_text.splitlines()
    first_content_seen = False
    for line in lines:
        stripped = line.strip()
        if stripped.startswith("Zadanie"):
            p = doc.add_paragraph()
            run = p.add_run(line)
            run.bold = True
        elif not first_content_seen and stripped:
            # Pierwsza niepusta linia to zwykle tytuł tematu — wyróżniamy jako nagłówek.
            doc.add_heading(line, level=1)
            first_content_seen = True
        elif stripped == "":
            doc.add_paragraph("")
        else:
            doc.add_paragraph(line)
            first_content_seen = True

    if key_text:
        key_lines = key_text.splitlines()
        # Pierwsza linia to nagłówek sekcji (np. "KLUCZ ODPOWIEDZI (dla nauczyciela)") —
        # zastępujemy ją własnym nagłówkiem Word, resztę wypisujemy jak jest.
        remaining = key_lines[1:] if key_lines and ANSWER_KEY_MARKER in key_lines[0] else key_lines
        doc.add_page_break()
        doc.add_heading("Klucz odpowiedzi (dla nauczyciela)", level=2)
        for line in remaining:
            doc.add_paragraph(line)

    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))


def save_as_txt(text: str, path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(text, encoding="utf-8")


def produce_one(
    client: OpenAI,
    params: KartaParams,
    model: str,
    output_base: Path,
) -> None:
    print(f"-> Generuję: {params.przedmiot} / klasa {params.klasa} / temat: {params.temat}")
    text = generate_worksheet_text(client, params, model)
    docx_path = output_base.with_suffix(".docx")
    txt_path = output_base.with_suffix(".txt")
    save_as_docx(text, params, docx_path)
    save_as_txt(text, txt_path)
    print(f"   Zapisano: {docx_path}")
    print(f"   Zapisano: {txt_path}")


def run_single(args: argparse.Namespace, client: OpenAI) -> None:
    params = KartaParams(
        przedmiot=args.przedmiot,
        klasa=str(args.klasa),
        temat=args.temat,
        wskazowki=args.wskazowki or "brak",
        liczba_zadan=str(args.liczba_zadan),
        trudnosc=args.trudnosc,
        typ_zadan=args.typ_zadan,
        klucz=args.klucz,
    )
    output_base = Path(args.output) if args.output else Path("karty") / slugify(params.temat)
    produce_one(client, params, args.model, output_base)


def run_batch(args: argparse.Namespace, client: OpenAI) -> None:
    csv_path = Path(args.batch)
    if not csv_path.exists():
        print(f"Nie znaleziono pliku CSV: {csv_path}", file=sys.stderr)
        sys.exit(1)

    output_dir = Path(args.output_dir or "karty")
    field_names = {f.name for f in fields(KartaParams)}
    errors: list[str] = []
    count = 0

    with csv_path.open(newline="", encoding="utf-8") as f:
        reader = csv.DictReader(f)
        for row_num, row in enumerate(reader, start=2):
            kwargs = {k: v for k, v in row.items() if k in field_names and v not in (None, "")}
            if "temat" not in kwargs or not kwargs["temat"].strip():
                errors.append(f"Wiersz {row_num}: brak wymaganej kolumny 'temat', pomijam.")
                continue
            params = KartaParams(**{**KartaParams().__dict__, **kwargs})
            output_base = output_dir / f"{count + 1:02d}_{slugify(params.temat)}"
            try:
                produce_one(client, params, args.model, output_base)
                count += 1
            except Exception as exc:  # noqa: BLE001
                errors.append(f"Wiersz {row_num} ('{params.temat}'): {exc}")

    print(f"\nGotowe: wygenerowano {count} kart(y) pracy w katalogu '{output_dir}'.")
    if errors:
        print(f"\nBłędy ({len(errors)}):", file=sys.stderr)
        for e in errors:
            print(f"  - {e}", file=sys.stderr)


def build_arg_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description="Generator Kart Pracy — produkcja kart pracy dla szkoły podstawowej."
    )
    parser.add_argument("--api-key", help="Klucz OpenAI API (domyślnie zmienna OPENAI_API_KEY).")
    parser.add_argument(
        "--model", default="gpt-4o-mini", help="Model OpenAI do użycia (domyślnie gpt-4o-mini)."
    )

    mode = parser.add_mutually_exclusive_group(required=True)
    mode.add_argument("--temat", help="Temat karty pracy (tryb pojedynczy).")
    mode.add_argument("--batch", help="Ścieżka do pliku CSV z listą tematów (tryb wsadowy).")

    parser.add_argument("--przedmiot", default="Matematyka", help="Przedmiot szkolny.")
    parser.add_argument("--klasa", default="4", help="Klasa (1-8).")
    parser.add_argument("--wskazowki", default="", help="Dodatkowe wskazówki dla GPT.")
    parser.add_argument("--liczba-zadan", default="5", help="Liczba zadań na karcie.")
    parser.add_argument(
        "--trudnosc",
        default="zróżnicowany (od łatwych do trudnych)",
        help="Poziom trudności zadań.",
    )
    parser.add_argument(
        "--typ-zadan",
        default="mieszane",
        help="Typy zadań (mieszane, zamknięte, otwarte, uzupełnianie luk, ...).",
    )
    parser.add_argument("--klucz", default="tak", choices=["tak", "nie"], help="Dołączyć klucz odpowiedzi?")

    parser.add_argument("--output", help="Ścieżka bazowa pliku wyjściowego (tryb pojedynczy, bez rozszerzenia).")
    parser.add_argument("--output-dir", help="Katalog wyjściowy (tryb wsadowy).")

    return parser


def main() -> None:
    parser = build_arg_parser()
    args = parser.parse_args()

    api_key = args.api_key or os.environ.get("OPENAI_API_KEY")
    if not api_key:
        print(
            "Brak klucza API. Ustaw zmienną OPENAI_API_KEY albo użyj --api-key.",
            file=sys.stderr,
        )
        sys.exit(1)

    client = OpenAI(api_key=api_key)

    if args.batch:
        run_batch(args, client)
    else:
        run_single(args, client)


if __name__ == "__main__":
    main()
